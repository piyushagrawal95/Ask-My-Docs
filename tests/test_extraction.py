import pytest
from io import BytesIO
import docx as docx_lib
from pypdf import PdfWriter

from app.services.extraction import extract_text,ExtractionError

def test_txt_extraction_returns_single_page():
    pages=extract_text(b"Hello World,this is a text","notes.txt")
    assert len(pages)==1
    assert pages[0][0]==1
    assert "Hello World" in pages[0][1]

def test_empty_txt_raises_extraction_error():
    with pytest.raises(ExtractionError):
        extract_text(b" ","empty.txt")

def test_unsupported_extension_raises_extraction_error():
    with pytest.raises(ExtractionError):
        extract_text(b"some bytes","notes.exe")

def test_docx_extraction_returns_text():
    buf=BytesIO()
    document=docx_lib.Document()
    document.add_paragraph("This is a paragraph from a Word document")
    document.save(buf)

    pages=extract_text(buf.getvalue(),"notes.docx")
    assert len(pages)==1
    assert "paragraph from a Word document" in pages[0][1]

def test_empty_docx_raises_extraction_error():
    buf=BytesIO()
    document=docx_lib.Document()
    document.save(buf)

    with pytest.raises(ExtractionError):
        extract_text(buf.getvalue(),"empty.docx")


def test_blank_pdf_raises_extraction_error():
    buf=BytesIO()
    writer=PdfWriter()
    writer.add_blank_page(width=200,height=200)
    writer.write(buf)

    with pytest.raises(ExtractionError):
        extract_text(buf.getvalue(),"scanned.pdf")


def test_docx_extraction_preserves_tables():
    buf = BytesIO()
    document = docx_lib.Document()
    document.add_paragraph("Introductory paragraph")
    
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Laptop"
    table.cell(1, 1).text = "$1200"
    table.cell(2, 0).text = "Phone"
    table.cell(2, 1).text = "$800"

    document.add_paragraph("Concluding remarks")
    document.save(buf)

    pages = extract_text(buf.getvalue(), "catalog.docx")
    assert len(pages) == 1
    content = pages[0][1]

    # Verify presence of text and table markers
    assert "Introductory paragraph" in content
    assert "<<<TABLE>>>" in content
    assert "| Product | Price |" in content
    assert "|---|---|" in content
    assert "| Laptop | $1200 |" in content
    assert "| Phone | $800 |" in content
    assert "<<<END_TABLE>>>" in content
    assert "Concluding remarks" in content

    # Verify natural reading order
    intro_idx = content.find("Introductory paragraph")
    table_idx = content.find("<<<TABLE>>>")
    conclusion_idx = content.find("Concluding remarks")
    assert intro_idx < table_idx < conclusion_idx


def test_docx_table_only_extraction():
    buf = BytesIO()
    document = docx_lib.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "98%"
    document.save(buf)

    pages = extract_text(buf.getvalue(), "metrics.docx")
    assert len(pages) == 1
    assert "Metric" in pages[0][1]
    assert "Accuracy" in pages[0][1]