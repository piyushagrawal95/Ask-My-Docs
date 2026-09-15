from app.services.chunking import chunk_pages

def test_chunk_pages_empty_input_returns_empty_list():
    assert chunk_pages([])==[]

def test_chunk_pages_preserve_page_number():
    pages=[(1,"Short text on page one"),(2,"Short text on page two")]
    chunk=chunk_pages(pages)
    assert len(chunk)==2
    assert chunk[0].page_number==1
    assert chunk[1].page_number==2


def test_chunk_pages_split_long_text():
    long_text=("This is a sentence."*200)
    chunks=chunk_pages([(1,long_text)])

    assert len(chunks)>1
    assert [c.chunk_index for c in chunks]==list(range(len(chunks)))
    assert all(c.page_number==1 for c in chunks)


def test_chunk_pages_skips_blank_pieces():
    chunks=chunk_pages([(1," \n\n    ")])
    assert chunks==[]


def test_chunk_pages_docx_table_preservation():
    from io import BytesIO
    import docx as docx_lib
    from app.services.extraction import extract_text

    buf = BytesIO()
    doc = docx_lib.Document()
    doc.add_paragraph("Table description")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "100"
    table.cell(1, 1).text = "200"
    table.cell(2, 0).text = "300"
    table.cell(2, 1).text = "400"
    doc.save(buf)

    pages = extract_text(buf.getvalue(), "data.docx")
    chunks = chunk_pages(pages)

    assert len(chunks) >= 1
    # Check that the table chunk contains markdown table structure and values
    all_chunks_text = "\n".join(c.content for c in chunks)
    assert "| Column A | Column B |" in all_chunks_text
    assert "| 100 | 200 |" in all_chunks_text
    assert "| 300 | 400 |" in all_chunks_text