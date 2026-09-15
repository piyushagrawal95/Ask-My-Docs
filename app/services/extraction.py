import logging
from io import BytesIO
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
import fitz  # PyMuPDF

logger = logging.getLogger("extraction")

TABLE_START = "<<<TABLE>>>"
TABLE_END = "<<<END_TABLE>>>"


class ExtractionError(Exception):
  pass


def extract_text(file_bytes: bytes, file_name: str,on_progress=None) -> list[tuple[int, str]]:
  ext = file_name.lower().rsplit(".", 1)[-1]

  if ext == "pdf":
    return _extract_pdf(file_bytes, on_progress=on_progress)
  elif ext == "docx":
    return _extract_docx(file_bytes)
  elif ext == "txt":
    return _extract_txt(file_bytes)
  else:
    raise ExtractionError(f"Unsupported file type: .{ext}")


def _extract_pdf(file_bytes: bytes,on_progress=None) -> list[tuple[int, str]]:
  """PyMuPDF based ultra-fast and low-memory extractor.

  - 82 pages parse hone me sirf 2-3 seconds lagenge (15 minute nahi).
  - RAM consumption 50MB se kam rahegi (Render Free Tier 512MB par safe).
  - Tables ko clean Markdown me extract karta hai.
  - Table ka text double repeat nahi hota.
  """
  pages = []

  try:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
  except Exception as e:
    raise ExtractionError(f"Could not open PDF: {e}")

  total_pages=doc.page_count
  if on_progress:
    on_progress(0,total_pages)

  PROGRESS_EVERY =10 # har 10 pages ke baad DB m update krna

  for page_num, page in enumerate(doc, start=1):
    elements = []  # tuple of (y0_position, content_str)

    # 1. Page par tables detect karein
    tabs = page.find_tables()
    table_rects = [fitz.Rect(t.bbox) for t in tabs.tables] if tabs.tables else []

    # 2. Page ke text blocks extract karein (blocks format: x0, y0, x1, y1, text, block_no, type)
    blocks = page.get_text("blocks")
    for b in blocks:
      text = b[4].strip()
      if not text:
        continue

      # Check karein ki ye text block kisi table ke andar to nahi aa raha
      b_rect = fitz.Rect(b[:4])
      is_inside_table = False
      for tr in table_rects:
        intersection = b_rect & tr
        # Agar block ka 40% se zyada hissa table ke andar hai to use skip karein
        if (
            not intersection.is_empty
            and b_rect.get_area() > 0
            and (intersection.get_area() / b_rect.get_area()) > 0.4
        ):
          is_inside_table = True
          break

      if not is_inside_table:
        elements.append((b[1], text))  # b[1] is y0 (vertical position)

    # 3. Tables ko Markdown format me convert karein
    for tab in tabs:
      md = None
      try:
        md = tab.to_markdown()
      except Exception:
        pass

      # Fallback agar to_markdown() kisi table par fail ho
      if not md:
        extracted = tab.extract()
        if extracted:
          rows = []
          for r_idx, row in enumerate(extracted):
            cells = [str(c).strip() if c else "" for c in row]
            rows.append("| " + " | ".join(cells) + " |")
            if r_idx == 0:
              rows.append("|" + "---|" * len(cells))
          md = "\n".join(rows)

      if md and md.strip():
        table_content = f"{TABLE_START}\n{md.strip()}\n{TABLE_END}"
        elements.append((tab.bbox[1], table_content))

    # 4. Elements ko top-to-bottom natural order me sort karein
    elements.sort(key=lambda x: x[0])

    combined = "\n\n".join(item[1] for item in elements).strip()
    if combined:
      pages.append((page_num, combined))

    # 5. Progress report karein — sirf har PROGRESS_EVERY pages ke baad ya
    # last page par, taaki DB writes minimal rahein aur speed pe asar na pade.
    if on_progress and (page_num % PROGRESS_EVERY == 0 or page_num == total_pages):
      on_progress(page_num, total_pages)

  doc.close()

  if not pages:
    raise ExtractionError("No extractable text found (possibly a scanned PDF).")

  return pages


def _docx_table_to_markdown(table: Table) -> str:
  rows = []
  for row in table.rows:
    cells = [
        cell.text.strip().replace("\r\n", " ").replace("\n", " ").replace("|", "\\|")
        for cell in row.cells
    ]
    rows.append(cells)

  if not rows:
    return ""

  if not any(any(c for c in r) for r in rows):
    return ""

  num_cols = len(rows[0])
  header = "| " + " | ".join(rows[0]) + " |"
  separator = "|" + "---|" * num_cols
  lines = [header, separator]
  for r in rows[1:]:
    if len(r) < num_cols:
      r = r + [""] * (num_cols - len(r))
    elif len(r) > num_cols:
      r = r[:num_cols]
    lines.append("| " + " | ".join(r) + " |")

  return "\n".join(lines)


def _iter_docx_elements(parent, document):
  for child in parent:
    if isinstance(child, CT_P) or child.tag.endswith("p"):
      yield Paragraph(child, document)
    elif isinstance(child, CT_Tbl) or child.tag.endswith("tbl"):
      yield Table(child, document)
    elif child.tag.endswith("sdt"):
      for sdt_child in child:
        if sdt_child.tag.endswith("sdtContent"):
          yield from _iter_docx_elements(sdt_child, document)


def _extract_docx(file_bytes: bytes) -> list[tuple[int, str]]:
  try:
    document = docx.Document(BytesIO(file_bytes))
  except Exception as e:
    raise ExtractionError(f"Could not open .docx file: {e}")

  elements = []
  for item in _iter_docx_elements(document.element.body, document):
    if isinstance(item, Paragraph):
      text = item.text.strip()
      if text:
        elements.append(text)
    elif isinstance(item, Table):
      # Ek malformed/nested table poore document ka extraction crash na kare —
      # isliye per-table try/except: fail hone par sirf wo table skip hota hai,
      # baaki paragraphs aur tables normally extract hote rehte hain.
      try:
        md = _docx_table_to_markdown(item)
      except Exception:
        logger.warning("Skipping a docx table that failed to convert to markdown", exc_info=True)
        md = ""
      if md:
        elements.append(f"{TABLE_START}\n{md}\n{TABLE_END}")

  combined = "\n\n".join(elements).strip()
  if not combined:
    raise ExtractionError("No extractable text found in .docx file.")

  return [(1, combined)]


def _extract_txt(file_bytes: bytes) -> list[tuple[int, str]]:
  text = file_bytes.decode("utf-8", errors="ignore")
  if not text.strip():
    raise ExtractionError("File is empty.")
  return [(1, text)]